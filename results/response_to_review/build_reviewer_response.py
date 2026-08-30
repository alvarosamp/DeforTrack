from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "outputs/reviewer_response/resposta_aos_revisores_defortrack.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(31, 78, 121)
    return p


def add_para(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Resposta aos revisores e resumo das modificações")
run.bold = True
run.font.size = Pt(16)
run.font.name = "Calibri"
run.font.color.rgb = RGBColor(31, 78, 121)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Artigo: DeforTrack: A High-Resolution Dataset for Deep Learning-Based Deforestation Segmentation")
r.italic = True
r.font.size = Pt(11)

add_heading(doc, "Carta de agradecimento", 1)
add_para(
    doc,
    "Prezados Revisores,",
)
add_para(
    doc,
    "Gostaríamos de agradecer sinceramente pelo tempo dedicado à leitura do manuscrito e pelos comentários detalhados. "
    "As observações recebidas foram fundamentais para aprimorar a clareza metodológica, a organização dos resultados, "
    "a discussão sobre generalização e a apresentação visual do artigo. Revisamos o manuscrito considerando cada ponto levantado "
    "e descrevemos abaixo as principais modificações realizadas.",
)
add_para(
    doc,
    "As alterações foram feitas com o objetivo de tornar o texto mais preciso, evitar afirmações não comprovadas, melhorar a rastreabilidade "
    "dos experimentos e apresentar de forma mais honesta as limitações da validação atual.",
)

add_heading(doc, "Resumo geral das modificações", 1)
bullets = [
    "A descrição do conjunto de dados foi revisada para remover afirmações específicas sobre fontes públicas que não deveriam ser explicitamente citadas no texto principal.",
    "A seção de Data Augmentation foi renomeada para Data Augmentation and Availability e recebeu o rótulo de referência cruzada correspondente.",
    "Foi adicionada uma explicação sobre o procedimento usado para medir RAM, CPU, tempo de inferência, FPS e consumo de energia.",
    "As tabelas de métricas foram revisadas para destacar em negrito os melhores valores de desempenho e eficiência computacional.",
    "O parágrafo de resultados referente aos modelos YOLOv11 foi corrigido para ficar consistente com os valores oficiais das tabelas.",
    "Foi adicionada uma avaliação externa com dois datasets adicionais, separando os resultados numéricos das imagens de exemplo para melhorar a legibilidade.",
    "A limitação relacionada à ausência de validação cross-dataset estrita foi explicitada de forma curta e honesta.",
    "As referências bibliográficas foram revisadas para incluir os datasets externos usados na nova avaliação.",
]
for item in bullets:
    add_bullet(doc, item)

add_heading(doc, "Resposta ponto a ponto", 1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
hdr = table.rows[0].cells
headers = ["Comentário / ponto revisado", "Modificação realizada", "Local no manuscrito"]
for cell, text in zip(hdr, headers):
    set_cell_text(cell, text, bold=True)
    set_cell_shading(cell, "F2F4F7")

rows = [
    (
        "Clareza sobre as fontes do dataset",
        "O texto foi ajustado para descrever as três fontes principais de forma mais geral: UAV, Google Earth Engine e repositórios públicos ambientais. Foram removidas menções específicas que não procediam no contexto final do artigo.",
        "Dataset Implementation - Data Collection",
    ),
    (
        "Disponibilidade do dataset e referência cruzada",
        "A subseção passou a se chamar Data Augmentation and Availability e foi adicionado o label sec:data_augmentation para corrigir a referência interna do texto.",
        "Dataset Implementation - Data Augmentation and Availability",
    ),
    (
        "Procedimento de medição computacional",
        "Foi acrescentado que RAM, CPU, tempo de inferência, FPS e energia foram medidos pelo mesmo procedimento para todos os modelos. Também foi informado o uso de psutil e NVML.",
        "Deep-learning Semantic Segmentation Models Evaluation - Evaluation Metrics",
    ),
    (
        "Consistência dos resultados YOLOv11",
        "O parágrafo de discussão dos YOLOv11 foi reescrito para refletir os valores oficiais das tabelas, destacando YOLOv11m para IoU, YOLOv11x para métricas de contorno e YOLOv11n para compacidade.",
        "Results and Discussion",
    ),
    (
        "Destaque dos melhores valores nas tabelas",
        "As melhores métricas de segmentação e eficiência computacional foram destacadas em negrito, incluindo IoU, mAP, Dice/F1, Boundary IoU/F1, Pixel Accuracy, Precision, Recall, tamanho, RAM, tempo, FPS, CPU e energia.",
        "Tables: Segmentation metrics and Computational metrics",
    ),
    (
        "Cross-dataset validation",
        "Foi esclarecido que a validação interna original não constitui validação cross-dataset estrita. Para responder ao comentário, foram realizados testes externos em dois datasets adicionais semanticamente compatíveis.",
        "Results and Discussion - External Cross-Dataset Evaluation",
    ),
    (
        "Apresentação dos datasets externos",
        "A tabela externa foi simplificada para conter apenas os resultados numéricos e interpretações. As imagens representativas foram movidas para uma figura separada, aumentando a legibilidade.",
        "Table: External cross-dataset evaluation; Figure: external dataset examples",
    ),
    (
        "Limitações e generalização",
        "Foi adicionada uma limitação curta explicando que a validação cross-dataset estrita permanece como trabalho futuro devido à ausência de metadados de origem por imagem no dataset exportado.",
        "Limitations and Generalization Considerations",
    ),
    (
        "Referências",
        "As referências foram revisadas para incluir os datasets externos usados na nova avaliação e para manter coerência entre as citações do texto e a bibliografia.",
        "References",
    ),
]

for point, modification, location in rows:
    cells = table.add_row().cells
    set_cell_text(cells[0], point)
    set_cell_text(cells[1], modification)
    set_cell_text(cells[2], location)

for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(2)

add_heading(doc, "Texto sugerido para acompanhar a submissão", 1)
add_para(
    doc,
    "We thank the reviewers for their careful reading of the manuscript and for their constructive comments. "
    "The manuscript was revised to improve the dataset description, clarify the evaluation protocol, correct inconsistencies in the reported results, "
    "and provide a more transparent discussion of generalization limitations. In particular, we revised the data collection description, added details about "
    "computational profiling, corrected the YOLOv11 discussion according to the final metrics, highlighted the best values in the tables, and added an external "
    "evaluation using two additional forest segmentation datasets. We also explicitly state that strict cross-dataset validation remains future work due to the "
    "absence of image-level source metadata in the final exported dataset.",
)

add_heading(doc, "Observação final", 1)
add_para(
    doc,
    "A versão revisada evita afirmar que os resultados internos demonstram generalização cross-dataset estrita. Em vez disso, o manuscrito apresenta os resultados "
    "internos como benchmark do DeforTrack e os novos testes externos como evidência complementar de generalização parcial sob mudança de domínio.",
)

add_heading(doc, "Complemento: pontos adicionais de revisao", 1)
add_para(
    doc,
    "Esta secao complementa a resposta aos revisores com os pontos adicionais recebidos durante o processo de revisao. "
    "As respostas abaixo indicam o que foi modificado no manuscrito, o que foi esclarecido como limitacao e o que permanece como trabalho futuro.",
)

extra = doc.add_table(rows=1, cols=4)
extra.alignment = WD_TABLE_ALIGNMENT.CENTER
extra.style = "Table Grid"
extra_headers = ["Revisor", "Comentario", "Resposta / modificacao", "Status"]
for cell, text in zip(extra.rows[0].cells, extra_headers):
    set_cell_text(cell, text, bold=True)
    set_cell_shading(cell, "F2F4F7")

extra_rows = [
    (
        "Reviewer 1",
        "Dataset creation process is common practice; superiority over existing datasets not justified.",
        "Foi adicionada uma tabela comparativa objetiva com numero de imagens/produtos, resolucao espacial, tipo de anotacao, classes, diversidade geografica e uso pretendido. A justificativa foi reformulada para evitar uma afirmacao generica de superioridade e destacar a contribuicao especifica do DeforTrack.",
        "Respondido parcialmente; reforcar com estatisticas quantitativas quando disponiveis.",
    ),
    (
        "Reviewer 1",
        "Global Forest Watch lacks sharpness at higher zoom levels should be referenced.",
        "A frase deve ser substituida por uma formulacao tecnicamente mais precisa: Global Forest Watch usa produtos Landsat de 30 m, que sao menos granulares para pequenos disturbios e limites locais finos. Essa afirmacao pode ser referenciada com material do GFW/WRI sobre resolucao de 30 m e limitacoes de granularidade.",
        "Ajustar texto e referencia.",
    ),
    (
        "Reviewer 1",
        "Novelty claim not rigorously supported.",
        "A novidade foi reforcada com a tabela comparativa e com a descricao objetiva dos elementos combinados no DeforTrack: imagens RGB 640 x 640, mascaras manuais por poligono, protocolo binario floresta/nao-floresta, diversidade de fontes e avaliacao orientada a desempenho e implantacao.",
        "Respondido; pode ser fortalecido com estatisticas de diversidade.",
    ),
    (
        "Reviewer 1",
        "Discussion remains superficial.",
        "A discussao foi expandida para incluir comportamento por familia de modelos, custo computacional, qualidade de fronteira, cenarios desafiadores e capacidade de generalizacao. Tambem foram adicionados casos provaveis de falha: bordas ambiguas, vegetacao esparsa, solo exposto, sombras, areas agricolas, vegetacao seca, nuvens e baixo contraste.",
        "Respondido no manuscrito.",
    ),
    (
        "Reviewer 1",
        "No inter-annotator agreement or labeling quality assessment.",
        "Como nao houve estudo formal de concordancia entre anotadores na versao atual, isso deve ser declarado como limitacao. Foi proposta metodologia futura: anotacao independente de uma amostra representativa, calculo de IoU/Dice entre anotadores e revisao por especialista dos casos de discordancia.",
        "Nao feito experimentalmente; responder como limitacao e trabalho futuro.",
    ),
    (
        "Reviewer 1",
        "Table I overwhelming; no variance/repeatability.",
        "Foi adicionada a formulacao de intervalo de confianca de 95% para as principais metricas por imagem. As tabelas tambem foram separadas entre metricas de segmentacao e metricas computacionais para melhorar legibilidade.",
        "Respondido parcialmente; incluir valores numericos de IC se disponiveis.",
    ),
    (
        "Reviewer 3",
        "Train/val/test split and leakage.",
        "O texto foi ajustado para esclarecer que o split foi feito antes da augmentacao, reduzindo vazamento direto por copias aumentadas. Ao mesmo tempo, foi reconhecido que a ausencia de metadados completos por imagem impede garantir split geografico, temporal ou por fonte. Os resultados devem ser interpretados como benchmark interno, nao como prova de generalizacao espacial/temporal estrita.",
        "Respondido como limitacao; futuro trabalho com metadados de fonte, regiao, bioma, campanha e data.",
    ),
]

for row in extra_rows:
    cells = extra.add_row().cells
    for cell, text in zip(cells, row):
        set_cell_text(cell, text)

add_heading(doc, "Trechos recomendados para o manuscrito", 1)
add_para(
    doc,
    "Para Global Forest Watch, substituir qualquer frase subjetiva como 'lacks sharpness at higher zoom levels' por uma frase tecnica e referenciada: "
    "'Global Forest Watch tree-cover products commonly rely on Landsat-based 30 m resolution data, which are less suitable for identifying small disturbances or fine-grained local forest boundaries.'",
)
add_para(
    doc,
    "Para qualidade de anotacao: 'Although all masks were manually annotated and reviewed, a formal inter-annotator agreement study was not performed in the current version. Future releases will include independent annotation of a representative subset of images and agreement analysis using mask-level IoU/Dice scores, followed by expert adjudication of disagreement cases.'",
)
add_para(
    doc,
    "Para leakage: 'Although the split was performed before data augmentation, the current exported dataset does not preserve complete image-level metadata regarding source, geographic region, acquisition date, or flight campaign. Therefore, strict spatial/temporal leakage analysis and region-wise validation could not be fully guaranteed. Future versions of DeforTrack will preserve this metadata to enable geographic, temporal, and source-aware validation protocols.'",
)

doc.save(OUT)
print(OUT)
