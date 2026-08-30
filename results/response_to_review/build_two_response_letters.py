from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = "outputs/reviewer_response"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_text(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def setup_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.italic = True
    r.font.size = Pt(10.5)
    return doc


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(31, 78, 121)
    return p


def para(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        set_text(cell, text, bold=True)
        shade(cell, "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            set_text(cell, text)
    return table


professor_rows = [
    (
        "Comparacao com datasets existentes",
        "Foi preparada uma tabela comparativa objetiva com numero de imagens/produtos, resolucao, tipo de anotacao, classes, diversidade geografica e uso pretendido.",
        "Introducao / Tabela comparativa",
    ),
    (
        "Afirmacao sobre Global Forest Watch",
        "A frase subjetiva sobre falta de nitidez deve ser substituida por uma formulacao tecnica: produtos GFW baseados em Landsat possuem resolucao de 30 m e sao menos adequados para pequenos disturbios e fronteiras locais finas.",
        "Introducao / comparacao com datasets",
    ),
    (
        "Novelty claim",
        "A novidade foi reformulada para nao depender de superioridade generica. O foco passou a ser a combinacao de imagens RGB de alta resolucao, mascaras manuais, protocolo floresta/nao-floresta e avaliacao orientada a implantacao.",
        "Introducao / contribuicoes",
    ),
    (
        "Discussao superficial",
        "A discussao foi ampliada com comportamento por familia de modelos, custo computacional, qualidade de borda, cenarios de falha e capacidade de generalizacao.",
        "Results and Discussion / Limitations",
    ),
    (
        "Casos de falha",
        "Foram descritos cenarios provaveis de erro: bordas ambiguas, vegetacao esparsa, solo exposto, sombras, areas agricolas, vegetacao seca, nuvens e baixo contraste.",
        "Limitations and Generalization Considerations",
    ),
    (
        "Inter-annotator agreement",
        "Como nao existe estudo formal nesta versao, a resposta deve assumir a limitacao. Foi proposta metodologia futura com anotacao independente de amostra representativa, IoU/Dice entre anotadores e adjudicacao por especialista.",
        "Limitations / Future work",
    ),
    (
        "Intervalo de confianca",
        "Foi adicionada a formulacao de intervalo de confianca de 95% para metricas por imagem. Se os valores numericos estiverem disponiveis, eles podem ser incluidos na tabela ou no texto.",
        "Evaluation Metrics",
    ),
    (
        "Train/val/test split e leakage",
        "Foi esclarecido que o split foi feito antes da augmentacao, mas que a falta de metadados completos impede comprovar independencia geografica/temporal estrita. Isso foi tratado como limitacao.",
        "Limitations and Generalization Considerations",
    ),
    (
        "Cross-dataset externo",
        "Foram testados dois datasets externos semanticamente compativeis e a apresentacao foi separada em tabela numerica e figura com exemplos.",
        "External Cross-Dataset Evaluation",
    ),
]

reviewer_rows = [
    (
        "Reviewer 1",
        "Dataset creation process is common practice; superiority over existing datasets not justified.",
        "We added an objective comparison table including number of images/products, spatial resolution, annotation type, classes/labels, geographic diversity, and intended use. We revised the text to avoid an unsupported broad superiority claim and to define DeforTrack as a task-specific benchmark for high-resolution forest/non-forest segmentation.",
    ),
    (
        "Reviewer 1",
        "The statement about Global Forest Watch requires a reference.",
        "We revised the wording to a technically supported statement: GFW tree-cover products commonly rely on Landsat-based 30 m data, which are less suitable for small disturbances and fine-grained local boundaries. The revised statement is supported by GFW/WRI documentation on tree-cover datasets and spatial resolution.",
    ),
    (
        "Reviewer 1",
        "Novelty claim not rigorously supported.",
        "We strengthened the novelty discussion by linking the contribution to measurable dataset characteristics: 6,094 original images, 14,648 images after augmentation, 640 x 640 RGB patches, manual polygon-based masks, binary forest/non-forest labels, multiple acquisition sources, and deployment-oriented benchmarking.",
    ),
    (
        "Reviewer 1",
        "Discussion remains superficial.",
        "We expanded the discussion to include model-family behavior, accuracy-efficiency trade-offs, boundary quality, computational constraints, external generalization, and likely failure scenarios such as ambiguous forest boundaries, sparse vegetation, exposed soil, shadows, clouds, agricultural regions, dry vegetation, and low-contrast scenes.",
    ),
    (
        "Reviewer 1",
        "No inter-annotator agreement or labeling quality assessment.",
        "We acknowledge that a formal inter-annotator agreement study was not performed in the current version. The limitation is now explicitly stated, and future work will include independent annotation of a representative subset, mask-level IoU/Dice agreement analysis, and expert adjudication of disagreement cases.",
    ),
    (
        "Reviewer 1",
        "Table I overwhelming; no variance/repeatability.",
        "We reorganized the results into clearer segmentation and computational metric tables and added the 95% confidence interval formulation for image-level metrics to improve repeatability reporting.",
    ),
    (
        "Reviewer 3",
        "Train/validation/test split and leakage.",
        "We clarified that the split was performed before data augmentation, reducing direct leakage from augmented copies. However, because the final exported dataset does not preserve complete image-level source, region, date, or flight-campaign metadata, strict geographic/temporal leakage analysis could not be fully guaranteed. This is now discussed as a limitation and future work.",
    ),
]


def build_professor():
    doc = setup_doc(
        "Carta-resposta ao professor",
        "Resumo das correcoes implementadas e pendencias do artigo DeforTrack",
    )
    heading(doc, "Prezado Professor Felipe,", 1)
    para(
        doc,
        "Obrigado pelos comentarios e direcionamentos durante a revisao do manuscrito. Abaixo organizo, de forma objetiva, o que foi ajustado no artigo, como cada ponto dos revisores foi tratado e quais aspectos ainda precisam ser assumidos como limitacao ou trabalho futuro.",
    )
    heading(doc, "Resumo das acoes realizadas", 1)
    for item in [
        "A descricao do dataset foi revisada para evitar afirmacoes nao comprovadas sobre superioridade e fontes especificas.",
        "Foi incluida uma comparacao objetiva com outros datasets/plataformas de sensoriamento remoto.",
        "A discussao dos resultados foi ampliada, incluindo cenarios de falha e generalizacao.",
        "Foi adicionada avaliacao externa com dois datasets semanticamente compativeis.",
        "A ausencia de inter-annotator agreement formal foi tratada como limitacao, com metodologia proposta para versoes futuras.",
        "A questao de leakage espacial/temporal foi discutida com honestidade, destacando que o split foi feito antes da augmentacao, mas sem garantia estrita por falta de metadados.",
    ]:
        bullet(doc, item)
    heading(doc, "Pontos tratados", 1)
    add_table(doc, ["Ponto", "Resposta / acao", "Local sugerido"], professor_rows)
    heading(doc, "Pendencias que dependem de decisao/orientacao", 1)
    for item in [
        "Definir se sera possivel realizar uma validacao formal das anotacoes com mais de um anotador.",
        "Decidir se os intervalos de confianca serao apresentados apenas metodologicamente ou tambem com valores numericos.",
        "Confirmar se ha metadados suficientes para algum agrupamento por fonte, regiao, campanha ou data. Se nao houver, manter como limitacao.",
    ]:
        bullet(doc, item)
    para(
        doc,
        "Atenciosamente,\nAlvaro Sampaio Careli",
    )
    path = f"{OUT_DIR}/response_letter_professor_defortrack.docx"
    doc.save(path)
    return path


def build_reviewers():
    doc = setup_doc(
        "Response Letter to the Reviewers",
        "DeforTrack: A High-Resolution Dataset for Deep Learning-Based Deforestation Segmentation",
    )
    heading(doc, "Dear Reviewers,", 1)
    para(
        doc,
        "We sincerely thank the reviewers for their careful reading of the manuscript and for the constructive comments. The manuscript was revised to improve the justification of the dataset contribution, clarify the evaluation protocol, expand the discussion of limitations and failure cases, and provide a more transparent treatment of generalization and data leakage risks.",
    )
    para(
        doc,
        "Below, we summarize the main changes made in response to each comment.",
    )
    heading(doc, "Point-by-point response", 1)
    add_table(doc, ["Reviewer", "Comment", "Response / manuscript change"], reviewer_rows)
    heading(doc, "Additional clarification on Global Forest Watch", 1)
    para(
        doc,
        "We revised the wording related to Global Forest Watch to avoid a subjective statement. Instead of stating that it 'lacks sharpness at higher zoom levels,' the manuscript now refers to the spatial-resolution limitation more precisely: Global Forest Watch tree-cover products commonly rely on Landsat-based 30 m resolution data, which are less suitable for identifying small disturbances or fine-grained local forest boundaries. This statement is supported by GFW/WRI documentation describing 30 m tree-cover products and explaining that 10 m products improve monitoring at small local scales.",
    )
    heading(doc, "Remaining limitations explicitly acknowledged", 1)
    for item in [
        "A formal inter-annotator agreement study was not performed in the current version.",
        "Strict geographic or temporal split validation could not be fully guaranteed because complete image-level source metadata were not preserved in the final exported dataset.",
        "The external datasets provide complementary evidence of generalization under domain shift, but they do not replace a fully source-aware cross-dataset validation protocol.",
    ]:
        bullet(doc, item)
    para(
        doc,
        "Sincerely,\nAlvaro Sampaio Careli",
    )
    path = f"{OUT_DIR}/response_letter_reviewers_defortrack.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_professor())
    print(build_reviewers())
